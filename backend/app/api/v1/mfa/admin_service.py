"""
MFA administration service layer.

Hosts the admin aggregation pipelines, Excel export, and MFA-reset business
logic that used to live inline in ``mfa/routes.py`` so that routes stay
focused on HTTP concerns (Dependency Inversion).
"""
import io
import logging
import re
from typing import Optional

from fastapi import HTTPException, status
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.api.v1.ai_providers.interfaces import IAIProviderRepository
from app.api.v1.ai_providers.repository import AIProviderRepository
from app.api.v1.mfa.interfaces import IMfaRepository
from app.api.v1.users.schemas import UserInDB

logger = logging.getLogger(__name__)


class MfaAdminService:
    """Service class handling MFA administration business logic."""

    def __init__(
        self,
        repository: IMfaRepository,
        ai_provider_repository: Optional[IAIProviderRepository] = None,
    ):
        """
        Initialize MFA admin service with its dependencies.

        Args:
            repository: MFA repository implementation
            ai_provider_repository: AI provider repository
                (defaults to AIProviderRepository)
        """
        self.repository = repository
        self.ai_provider_repository = (
            ai_provider_repository
            if ai_provider_repository is not None
            else AIProviderRepository()
        )

    async def list_users(self, page: int, page_size: int, search: str) -> dict:
        """List users with MFA status and usage aggregations.

        Supports pagination and case-insensitive search on email and
        full name. Aggregates project/diagram counts and connected AI
        models per user.

        Args:
            page: 1-based page number.
            page_size: Number of users per page.
            search: Optional substring filter for email or full name.

        Returns:
            Dict with ``items``, ``total``, ``page``, ``page_size``,
            and ``total_pages``.
        """
        # Build query
        query = UserInDB.find()
        if search:
            pattern = re.compile(re.escape(search), re.IGNORECASE)
            query = UserInDB.find(
                {"$or": [{"email": pattern}, {"full_name": pattern}]}
            )

        total = await query.count()

        skip = (page - 1) * page_size
        users = await query.skip(skip).limit(page_size).sort("-created_at").to_list()

        project_counts_by_user: dict[str, int] = {}
        diagram_counts_by_user: dict[str, int] = {}
        diagram_type_counts_by_user: dict[str, dict[str, int]] = {}
        try:
            from app.api.v1.projects.schemas import ProjectInDB
            from app.api.v1.diagrams.schemas import DiagramInDB

            user_ids = [str(user.id) for user in users]
            projects = await ProjectInDB.find({"user_id": {"$in": user_ids}}).to_list()
            project_owner_by_id = {str(project.id): project.user_id for project in projects}

            for project in projects:
                project_counts_by_user[project.user_id] = (
                    project_counts_by_user.get(project.user_id, 0) + 1
                )

            if project_owner_by_id:
                pipeline = [
                    {"$match": {"project_id": {"$in": list(project_owner_by_id)}}},
                    {
                        "$group": {
                            "_id": {"project_id": "$project_id", "diagram_type": "$diagram_type"},
                            "count": {"$sum": 1},
                        }
                    },
                ]
                grouped_diagrams = await DiagramInDB.get_motor_collection().aggregate(
                    pipeline
                ).to_list(None)

                for group in grouped_diagrams:
                    project_id = group["_id"]["project_id"]
                    user_id = project_owner_by_id[project_id]
                    diagram_type = str(group["_id"].get("diagram_type") or "unknown").lower()
                    count = group["count"]
                    diagram_counts_by_user[user_id] = (
                        diagram_counts_by_user.get(user_id, 0) + count
                    )
                    type_counts = diagram_type_counts_by_user.setdefault(user_id, {})
                    type_counts[diagram_type] = type_counts.get(diagram_type, 0) + count
        except Exception:
            logger.exception("Failed to aggregate admin user diagram counts")

        items = []
        for u in users:
            unused_codes = sum(
                1 for c in u.recovery_codes
                if not (c.used if hasattr(c, "used") else c.get("used", False))
            )
            last_login_at = u.last_login_at.isoformat() if u.last_login_at else None

            # Fetch subscription & plan name
            plan_name = None
            try:
                from app.api.v1.subscriptions.schemas import SubscriptionInDB, PlanInDB

                sub = await SubscriptionInDB.find_one(
                    SubscriptionInDB.user_id == str(u.id),
                    SubscriptionInDB.status == "active",
                )
                if sub:
                    plan = await PlanInDB.get(sub.plan_id)
                    if plan:
                        plan_name = plan.name
            except Exception:
                pass

            connected_ai_models = []
            try:
                ai_settings = await self.ai_provider_repository.get_user_settings(
                    str(u.id)
                )
                if ai_settings and ai_settings.providers:
                    for provider in ai_settings.providers:
                        if not provider.is_active:
                            continue
                        provider_name = (
                            provider.provider.value
                            if hasattr(provider.provider, "value")
                            else provider.provider
                        )
                        connected_ai_models.append(
                            {
                                "provider": provider_name,
                                "model": provider.model,
                                "is_default": provider.is_default,
                                "is_active": provider.is_active,
                                "display_name": provider.display_name,
                            }
                        )
            except Exception:
                pass

            user_id = str(u.id)
            project_count = project_counts_by_user.get(user_id, 0)
            diagram_count = diagram_counts_by_user.get(user_id, 0)
            diagram_type_counts = diagram_type_counts_by_user.get(user_id, {})

            items.append({
                "id": str(u.id),
                "email": u.email,
                "full_name": u.full_name,
                "role": u.role,
                "is_active": u.is_active,
                "mfa_enabled": u.mfa_enabled,
                "mfa_methods": u.mfa_methods,
                "mfa_default_method": u.mfa_default_method,
                "recovery_codes_remaining": unused_codes,
                "created_at": u.created_at.isoformat() if u.created_at else None,
                "plan_name": plan_name,
                "project_count": project_count,
                "diagram_count": diagram_count,
                "diagram_type_counts": diagram_type_counts,
                "connected_ai_models": connected_ai_models,
                "last_login_at": last_login_at,
            })

        return {
            "items": items,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size,
        }

    async def reset_user_mfa(self, user_id: str, admin_email: str) -> dict:
        """Reset (disable) all MFA methods for a user.

        Args:
            user_id: ID of the target user.
            admin_email: Email of the admin performing the reset.

        Returns:
            Dict with a user-facing ``message``.

        Raises:
            HTTPException 404: If the target user does not exist.
        """
        from bson import ObjectId

        target_user = await UserInDB.get(ObjectId(user_id))
        if not target_user:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="Usuario no encontrado",
            )

        if not target_user.mfa_enabled:
            return {"message": "El usuario no tiene MFA habilitado"}

        # Disable all methods
        for method in list(target_user.mfa_methods):
            await self.repository.disable_mfa(user_id, method)

        logger.info(
            "Admin %s reset MFA for user %s (%s)",
            admin_email,
            target_user.email,
            user_id,
        )

        # Audit log
        from app.api.v1.users.audit_log import (
            EVENT_ADMIN_MFA_RESET,
            log_event,
        )

        await log_event(
            EVENT_ADMIN_MFA_RESET,
            target_user.email,
            user_id=user_id,
            details=f"reset_by={admin_email}",
        )

        return {"message": f"MFA desactivado para {target_user.email}"}

    async def export_users_excel(self, lang: str) -> io.BytesIO:
        """Build the admin users Excel export.

        Args:
            lang: ``"en"`` or ``"es"`` selects the header language.

        Returns:
            An in-memory workbook buffer positioned at byte 0.
        """
        users = await UserInDB.find_all().sort("-created_at").to_list()

        wb = Workbook()
        ws = wb.active
        ws.title = "Users" if lang == "en" else "Usuarios"

        # Headers
        if lang == "en":
            headers = [
                "Email", "Full Name", "Role", "Plan", "License Usage", "Active",
                "MFA Enabled", "MFA Methods", "Default Method",
                "Recovery Codes Remaining", "Created At",
            ]
        else:
            headers = [
                "Correo", "Nombre completo", "Rol", "Plan", "Uso de licencia",
                "Activo", "MFA Habilitado", "Métodos MFA",
                "Método predeterminado", "Códigos de recuperación",
                "Fecha de registro",
            ]

        # Style
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(
            start_color="7C3AED", end_color="7C3AED", fill_type="solid"
        )
        header_alignment = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            bottom=Side(style="thin", color="E5E7EB"),
        )

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Data rows
        for row_idx, u in enumerate(users, 2):
            unused_codes = sum(
                1 for c in u.recovery_codes
                if not (c.used if hasattr(c, "used") else c.get("used", False))
            )
            methods_str = ", ".join(
                ("Email" if m == "email" else "TOTP") for m in u.mfa_methods
            )
            default_str = ""
            if u.mfa_default_method:
                default_str = "Email" if u.mfa_default_method == "email" else "TOTP"
            created_str = u.created_at.strftime("%Y-%m-%d %H:%M") if u.created_at else ""

            # Fetch plan name
            plan_name = ""
            try:
                from app.api.v1.subscriptions.schemas import SubscriptionInDB, PlanInDB

                sub = await SubscriptionInDB.find_one(
                    SubscriptionInDB.user_id == str(u.id),
                    SubscriptionInDB.status == "active",
                )
                if sub:
                    plan = await PlanInDB.get(sub.plan_id)
                    if plan:
                        plan_name = plan.name
            except Exception:
                pass

            # Count projects and diagrams
            project_count = 0
            diagram_count = 0
            try:
                from app.api.v1.projects.schemas import ProjectInDB
                from app.api.v1.diagrams.schemas import DiagramInDB

                user_projects = await ProjectInDB.find(
                    ProjectInDB.user_id == str(u.id)
                ).to_list()
                if user_projects:
                    project_count = len(user_projects)
                    project_ids = [str(p.id) for p in user_projects]
                    diagram_count = await DiagramInDB.find(
                        {"project_id": {"$in": project_ids}}
                    ).count()
            except Exception:
                pass

            yes = "Yes" if lang == "en" else "Sí"
            no = "No"
            project_label = "project" if project_count == 1 else "projects"
            diagram_label = "diagram" if diagram_count == 1 else "diagrams"
            if lang != "en":
                project_label = "proyecto" if project_count == 1 else "proyectos"
                diagram_label = "diagrama" if diagram_count == 1 else "diagramas"

            row_data = [
                u.email,
                u.full_name or "",
                u.role,
                plan_name,
                f"{project_count} {project_label} / {diagram_count} {diagram_label}",
                yes if u.is_active else no,
                yes if u.mfa_enabled else no,
                methods_str,
                default_str,
                unused_codes,
                created_str,
            ]
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = thin_border

        # Auto-width columns
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

        # Write to buffer
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        return buffer
